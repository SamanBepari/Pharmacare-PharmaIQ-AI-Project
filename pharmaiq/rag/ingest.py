"""
Document ingestion: turn an uploaded file into [(page, section, text)] blocks.

Supported: .pdf (PyMuPDF, falls back to pypdf), .docx (python-docx), .txt, .md
Heavy parsers are imported lazily so the app still starts if one is missing.
"""
import hashlib
import io
import re

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class IngestError(Exception):
    pass


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_blocks(filename: str, data: bytes):
    """Return list of dicts {page, section, text}."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED:
        raise IngestError(f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(SUPPORTED))}.")
    if not data:
        raise IngestError("The file is empty.")
    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    return _plain(data, markdown=(ext == ".md"))


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _plain(data: bytes, markdown=False):
    text = _decode(data)
    blocks, section, buf = [], "", []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            blocks.append({"page": None, "section": section, "text": body})
        buf.clear()

    for line in text.splitlines():
        m = re.match(r"^\s{0,3}(#{1,6})\s+(.*)$", line) if markdown else None
        if m:
            flush()
            section = m.group(2).strip()
        else:
            buf.append(line)
    flush()
    if not blocks:
        raise IngestError("No readable text found in the file.")
    return blocks


def _pdf(data: bytes):
    blocks = []
    try:
        import fitz  # PyMuPDF
        with fitz.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, 1):
                t = page.get_text("text").strip()
                if t:
                    blocks.append({"page": i, "section": "", "text": t})
    except ImportError:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise IngestError("PDF support needs PyMuPDF or pypdf: pip install pymupdf")
        reader = PdfReader(io.BytesIO(data))
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                blocks.append({"page": i, "section": "", "text": t})
    except Exception as exc:  # corrupt PDFs
        raise IngestError(f"Could not read the PDF: {exc}")
    if not blocks:
        raise IngestError("No selectable text found. Scanned PDFs need OCR before upload.")
    return blocks


def _docx(data: bytes):
    try:
        import docx
    except ImportError:
        raise IngestError("DOCX support needs python-docx: pip install python-docx")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestError(f"Could not read the DOCX file: {exc}")
    blocks, section, buf = [], "", []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            blocks.append({"page": None, "section": section, "text": body})
        buf.clear()

    for para in document.paragraphs:
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            section = para.text.strip()
        elif para.text.strip():
            buf.append(para.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                buf.append(" | ".join(cells))
    flush()
    if not blocks:
        raise IngestError("No readable text found in the document.")
    return blocks
