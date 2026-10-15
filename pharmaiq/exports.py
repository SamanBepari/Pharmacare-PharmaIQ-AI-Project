"""Downloadable reports built from stored result sets (real database rows only)."""
import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="1a6b3c")


def _safe(v):
    """Neutralise spreadsheet formula injection (cells starting with = + - @)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def to_xlsx(payload: dict) -> bytes:
    wb = Workbook()
    about = wb.active
    about.title = "About"
    rows = [("PharmaIQ report", ""), ("Title", payload.get("title", "")), ("Question", payload.get("question", "")),
            ("Tool / route", payload.get("tool", "")), ("Data as of", payload.get("as_of", "")), ("", ""),
            ("Note", "All figures come from the pharmacy database via validated read-only queries."),
            ("", "")]
    rows += [("Assumption", a) for a in payload.get("assumptions", [])]
    rows += [("Fact", f) for f in payload.get("facts", [])]
    for i, t in enumerate(payload.get("tables", [])):
        rows.append((f"Query {i + 1}", t.get("sql", "")))
    for r in rows:
        about.append([_safe(r[0]), _safe(r[1])])
    about["A1"].font = Font(bold=True, size=14)
    about.column_dimensions["A"].width = 16
    about.column_dimensions["B"].width = 100
    for i, t in enumerate(payload.get("tables", []), 1):
        name = "".join(c for c in (t.get("title") or f"Table {i}") if c not in "[]:*?/\\")[:28] or f"Table {i}"
        ws = wb.create_sheet(f"{i}-{name}"[:31])
        for c, h in enumerate(t["columns"], 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), HEADER_FILL, Alignment(horizontal="center")
        for r, row in enumerate(t["rows"], 2):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=_safe(v))
        for c in range(1, len(t["columns"]) + 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(payload: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    t = (payload.get("tables") or [{"columns": [], "rows": []}])[0]
    w.writerow(t["columns"])
    for row in t["rows"]:
        w.writerow([_safe(v) for v in row])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")
